"""Build the editable paper from its reviewed Markdown source."""
from pathlib import Path
import sys

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def hyperlink(paragraph, url):
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = url
    run.append(text)
    element.append(run)
    paragraph._p.append(element)


def build(output):
    doc = Document()
    section = doc.sections[0]
    section.page_width,section.page_height = Inches(8.5),Inches(11)
    section.top_margin,section.bottom_margin = Inches(.65),Inches(.65)
    section.left_margin,section.right_margin = Inches(.75),Inches(.75)
    for name in ("Normal","Title","Subtitle","Heading 1","Heading 2"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0,0,0)
    # The bundled default can contain an inherited title rule. Keep this paper
    # entirely typographic, including after conversion to PDF.
    for element in list(doc.styles.element.iter(qn("w:pBdr"))):
        element.getparent().remove(element)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    doc.styles["Title"].font.size = Pt(25)
    doc.styles["Title"].paragraph_format.space_after = Pt(9)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"),"PAGE")
    footer._p.append(field)
    lines = Path(__file__).with_name("cone_display_improvement_paper.md").read_text().splitlines()
    i = 0
    while i<len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line=="[PAGE]":
            doc.add_page_break()
        elif line.startswith("# "):
            doc.add_paragraph(line[2:],"Title")
        elif line.startswith("## "):
            doc.add_paragraph(line[3:],"Heading 1")
        elif line.startswith("https://"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            hyperlink(p,line)
        elif line.startswith("| "):
            rows = []
            while i<len(lines) and lines[i].startswith("| "):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(c.startswith("---") for c in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=0,cols=3)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            for column,width in zip(table.columns,(2.05,2.1,2.85)):
                column.width = Inches(width)
            for n,row in enumerate(rows):
                cells = table.add_row().cells
                for cell,value in zip(cells,row):
                    cell.text = value
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    properties = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement("w:tcBorders")
                    for side in ("top","left","bottom","right"):
                        border = OxmlElement("w:"+side)
                        for key,val in (("val","single"),("sz","4"),("color","D9D9D9")):
                            border.set(qn("w:"+key),val)
                        borders.append(border)
                    properties.append(borders)
                    margins = OxmlElement("w:tcMar")
                    for side in ("top","left","bottom","right"):
                        margin = OxmlElement("w:"+side)
                        margin.set(qn("w:w"),"100")
                        margin.set(qn("w:type"),"dxa")
                        margins.append(margin)
                    properties.append(margins)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(0)
                        p.paragraph_format.line_spacing = 1
                        for run in p.runs:
                            run.font.size = Pt(10.5)
                            run.bold = n==0
                    if n==0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"),"EDEDED")
                        properties.append(shading)
                if n==0:
                    repeated = OxmlElement("w:tblHeader")
                    table.rows[-1]._tr.get_or_add_trPr().append(repeated)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue
        else:
            doc.add_paragraph(line)
        i += 1
    doc.core_properties.title = "Improving the Live Circular Cone Display"
    doc.core_properties.subject = "Optical limitations and development options"
    doc.core_properties.author = "Pepper's Cone project"
    output = Path(output)
    output.parent.mkdir(parents=True,exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    build(sys.argv[1])
